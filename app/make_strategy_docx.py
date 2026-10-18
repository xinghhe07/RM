# -*- coding: utf-8 -*-
"""生成《RMUL 2027 哨兵 · 策略方案（当前版）》.docx

可重复运行，覆盖生成同名文件。内容更新直接改本脚本后重跑，
或在 Word 里直接编辑生成的文档（二者取其一，避免双向漂移）。

运行： python make_strategy_docx.py
"""
from __future__ import annotations

import sys

from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

OUT = r"C:\RM\docs\哨兵策略方案_当前版_20261016.docx"

EA = "微软雅黑"   # 中文字体
LATIN = "Calibri"


# --------------------------------------------------------------------------
# 样式助手
# --------------------------------------------------------------------------

def _set_style_font(style, size: float | None = None, bold: bool | None = None) -> None:
    style.font.name = LATIN
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), EA)
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold


def add_meta(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_bullets(doc: Document, items: list[str], level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    for it in items:
        doc.add_paragraph(it, style=style)


def add_table(doc: Document, rows: list[list[str]], widths: list[float] | None = None) -> None:
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = t.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(cell_text)
            run.font.size = Pt(9.5)
            if i == 0:
                run.font.bold = True
    doc.add_paragraph()


# --------------------------------------------------------------------------
# 文档内容
# --------------------------------------------------------------------------

def build() -> Document:
    doc = Document()

    # 全局字体
    _set_style_font(doc.styles["Normal"], size=10.5)
    for name, size in (("Title", 20), ("Heading 1", 14), ("Heading 2", 12)):
        try:
            _set_style_font(doc.styles[name], size=size)
        except KeyError:
            pass

    doc.add_heading("RMUL 2027 哨兵 · 策略方案（当前版）", level=0)
    add_meta(doc,
             "生成 2026-10-16 ｜ 状态：等待实车参数与官方答复后更新 ｜ 编制：视觉 + RL 组 ｜ "
             "配套代码：rm_sentry/tactical/（TacticalState、Rule、Safety、状态机接口）")

    # ---- 0 ----
    doc.add_heading("0. 一句话定位", level=1)
    doc.add_paragraph(
        "全向轮 + 前后双 Mid-360 的【地面】哨兵：不端到端、不机间通信、不上高地、压低过隧道；"
        "分层决策（规则 → 规划 → PPO 逐级打擂台），以控制度（CD）为统一目标货币；"
        "RL 不是生命线，Rule / Planner 保证不依赖 RL 也能参赛。")

    # ---- 1 ----
    doc.add_heading("1. 平台与约束（硬件 / 规则事实）", level=1)
    add_bullets(doc, [
        "底盘：全向轮 4 轮 X 型（第一年不做轮腿）。vx/vy/ω 三轴独立 → 横移、原地旋转、"
        "“边退边打”原生成立。",
        "感知：前后双 Mid-360（360° 几何 / 定位，自转不变）+ 自瞄相机（身份 / 瞄准）+ "
        "官方图传相机 + 裁判系统。",
        "规则：全自动运行、禁止机间通信、无官方定位数据 → 队友只能“看见”，不能“问到”。",
        "机动：地面单位。不上高地（第一版策略域设定，非永久架构原则）；压低车身过隧道"
        "（可行性待机械实测 + 官方答疑）。",
        "功率：底盘上限 100W + 超级电容缓冲，超限扣血 → 机动有“电费”，"
        "激进度必须预算化。",
    ])

    # ---- 2 ----
    doc.add_heading("2. 分工与接口边界", level=1)
    doc.add_paragraph("本文档范围：战术决策层（TacticalState、Rule / Planner / PPO、CD 经济学）"
                      "与下列接口约定。")
    add_table(doc, [
        ["模块", "负责", "接口约定"],
        ["自瞄", "其他同学", "视觉 → 战术层：目标身份（id / team / type / position / velocity / "
                             "confidence / age）——即 EnemyState 契约"],
        ["雷达 / 导航调试", "其他同学", "导航 → 战术层：path_cost / reachable / planner_ok；"
                                        "战术层不自己算路径代价"],
        ["状态机", "其他同学", "战术层 → 状态机：TacticalAction + point_id + reason（可日志化）"],
        ["仿真环境（含场地几何模型）", "视觉 + RL 组", "训练与验证在仿真内自闭环（fake executor），"
                                                 "不依赖其他组进度；场地几何 / 机器人运动学未到位时用占位并标注"],
        ["战术决策（本文档）", "视觉 + RL 组", "输入 = TacticalState 快照；输出 = 战术动作；"
                                             "与真机 / 仿真共用同一份代码（adapter 为接缝）"],
    ])

    # ---- 3 ----
    doc.add_heading("3. 感知与世界模型", level=1)
    add_bullets(doc, [
        "分工原则：雷达给几何（在哪），视觉给身份（是什么）。lidar-first，"
        "云台按需交叉提示给目标贴标签。",
        "信息边界（核心原则，写入 _core.py 头部并用测试逐字段守）：",
    ])
    add_bullets(doc, [
        "“通信不可得 ≠ 信息不可得”——所有进入 TacticalState 的非自身信息，"
        "必须能由哨兵自身合法传感器观测或由已有观测推导，并携带 age_s / confidence。",
        "A 类：可直读（自身 HP / 热量 / 弹量 / 读条 / 自身 buff）；"
        "B 类：可靠估计（敌人位置、机型、HP 档、控制度差值）；"
        "C 类：禁入（队友 HP / 任务 / 意图）。",
        "已落实的守卫：“不可见敌人不得进快照”（test_tactical_adapter.py）；"
        "后续扩为逐字段对抗性断言。",
    ])
    doc.add_paragraph("TacticalState 的 2027 增补（对现有 _core.py 的 diff）：")
    add_bullets(doc, [
        "比赛账本：our_cd / enemy_cd / cd_diff（裁判直读待查，或按“每控制 1 秒 +1、"
        "战亡 ±10”自估——CD 是低带宽积分器，可重建）。",
        "区域：双进度 our / enemy_progress（0~5）+ 在场 + 衰减动力学（无人激活双方 −1/s）"
        "+ time_to_flip；owner 降为派生字段（旧测试不破）。",
        "自机：+ buffer_fraction（功率缓冲）；虚弱改为真实语义（直到接触己方补给区）。",
        "敌人：+ type（装甲 ID）、HP 档（灯条粗读 + 伤害积分，双信道）、velocity / confidence。",
        "队友：独立 TeammateEstimate（只有位置 / 速度 / age——没有 HP / 任务）。",
        "增益：tunnel_buff（1.5× 热量冷却 60s）；高地增益随“不上高地”暂缓。",
        "比赛：remaining_time + phase（派生）+ score_gap。",
        "感知噪声注入（训练环境）：延迟 / 丢帧 / 误检 + age 联动，接在 adapter 缝上；"
        "actor 只看污染后的快照，critic 保留真值（非对称 actor-critic）。",
    ])

    # ---- 4 ----
    doc.add_heading("4. 决策架构（三层打擂台）", level=1)
    doc.add_paragraph("Rule V1（评分） → Planner V1.5（滚动时域规划） → PPO V2（战术选择）"
                      "——共用目标函数（期望 ΔCD）、共用信息集（真实可得，不许偷看真值）。")
    add_bullets(doc, [
        "硬门槛：PPO 打不过 Planner 不进比赛主线（Planner 代表“显式规划的上限”）。",
        "安全链：SafetyMonitor → 状态机 → Nav2 → 全向底盘。",
        "优先栈（不变）：安全 → 复活 / 虚弱 / 保护 → 导航有效 → 当前任务 → 重选。",
        "动作空间：战术任务 × 点 ID + mask（维持现状，不加底层动作）。",
        "运动模式（orbit / 游走 / 脱战机动）放执行层，不进 RL 动作空间"
        "（时标：战术 1~5s vs 运动 10~50ms）。",
        "赛前角色参数（holder / hunter / balancer）——待裁判答疑确认合法性。",
    ])

    # ---- 5 ----
    doc.add_heading("5. CD 经济学（统一目标货币）", level=1)
    add_bullets(doc, [
        "战亡 1 次 = −10 CD = 10 个“占点秒”（无人争夺下占 1 处 10 秒的收益）。",
        "独占 3 处 = 3/s；600 / 3 = 200s 清空对方 → 平均需 ≥2 处优势才能在 300s 内清零；"
        "1 处优势打满全场不够。",
        "回血 25%/s（补给区基础）→ 从 25% 回满约 3 秒：循环成本主要在路程，不在回血。",
        "死亡代价 ≈ 读条 5/10/15s + 复活 20% 血 + 虚弱（须回补给区解除）"
        "≈ 20~30s 非贡献时间 + −10 CD。",
        "撤退循环 ≈ 2×路程 + 3s，无 CD 损失 → 只要“能跑掉”，撤退经济上碾压死亡。",
        "差值 ≥200 / ≥400：劣势方读条 −33% / −50%、回血 33% / 50% → 落后时死亡更便宜"
        "（追赶机制）；撤退阈值不是固定值，HP* = f(ΔCD, 剩余时间, 区域价值, 敌情…)，由实验裁决。",
        "注：金币对哨兵无用（兑换发弹量仅限重装 / 步兵），不进哨兵状态。",
    ])

    # ---- 6 ----
    doc.add_heading("6. 哨兵人设（假设，实验裁决）", level=1)
    doc.add_paragraph(
        "“高速游击型驻点哨兵”：HP 是可循环资源，核心指标 = 单位时间创造的有效控制度"
        "（ΔCD/s），同时尽量不死亡。")
    doc.add_paragraph(
        "裁决实验（CD 曲线）：Holder / Hunter / Pinner / Quicksilver（机动守点 + 撤退循环）/ "
        "Hybrid 五种脚本策略，输出 ΔCD/s、存活占空比、占区时间、循环耗时等指标，"
        "多种子配对评估 → 用数据决定人设与资源投入。")

    # ---- 7 ----
    doc.add_heading("7. 推进方案顺序（三阶段，不跳）", level=1)

    doc.add_heading("阶段 A · 抽象战术层（现在就能做，只依赖“配置数字”，不依赖任何模型）", level=2)
    add_table(doc, [
        ["#", "做什么", "依赖", "出口（验收）"],
        ["A1", "2027 机制迁移：赛季档案 PPT → 官方手册；3 控制区双进度 / 控制度 600 / 战亡 −10 / "
               "读条 5+5 / 无敌 30 / 虚弱 / 补给；真区坐标进 config（依 docs/2027场地规格.md）",
         "无", "rules_test 2027 分支全过 + test_respawn 扩展过 + 跑完整一局"],
        ["A2", "日志闭环：JSONL（评分分量 + requested/executed/rejected + override 理由 + "
               "感知帧时间戳）", "A1", "跑一局可复盘每个决策"],
        ["A3", "CD 曲线实验：Holder / Hunter / Pinner / Quicksilver / Hybrid × 多种子配对",
         "A2", "ΔCD/s 表 → “快银是否成立”人设判决"],
        ["A4", "合成噪声层（延迟 / 丢帧 / 误检 + age 联动，接 adapter 缝，参数 [待标定]）"
               "+ CD 曲线带噪复跑", "A3", "stale / recovery 分支首次真实触发；人设结论不翻转"],
        ["A5", "Planner：滚动时域（区域动力学前滚 + P(death) / P(区易手)）", "A4",
         "稳定超过最好脚本策略"],
        ["A6", "PPO：噪声 + 对手池下训练（curriculum 干净→加噪；BC 热启动可选）", "A5",
         "稳定超过 Planner（多种子 + CI）→ 进比赛主线上线"],
    ])

    doc.add_heading("阶段 B · 保真升级（等模型）", level=2)
    add_table(doc, [
        ["#", "做什么", "等谁"],
        ["B1", "场地几何模型（归口：战术组）：图纸 → 2D / 2.5D 几何（区 / 墙 / 掩体 / 隧道 / 高程），"
               "全体标 [待实测]", "数据已备，模型待建"],
        ["B2", "机器人运动学：轮径 / 轴距 / v_max / a_max → 替换占位参数", "机械数单"],
        ["B3", "合流重基准（R2d）：替换占位几何 / 运动学 → CD 曲线与基线全部重跑；"
               "噪声参数由合成改实测标定", "B1 + B2"],
    ])

    doc.add_heading("阶段 C · 部署", level=2)
    add_table(doc, [
        ["#", "做什么", "出口"],
        ["C1", "ROS2 adapter 接通：真机话题 → TacticalState 逐字段对齐仿真",
         "同输入下仿真 / 真机决策一致"],
        ["C2", "真车联调 + 赛前彩排", "能跑完整局"],
    ])

    add_bullets(doc, [
        "外部并行（不阻塞 A）：机械数单 · 隧道净空 / 通行方式 · 交互卡检测范围 · "
        "《2027 通信协议》数据字典 · 赛前角色合法性答疑。",
        "水印纪律：阶段 A 全部数字带“占位几何 + 占位运动学”水印，B3 后作废重跑。",
        "验收先行：出口条件写死，没有出口不算完成。",
        "已有资产：TacticalState ✅ / 环境骨架 ✅ / Rule 基线 V1 ✅——实际新增从 A1 算起。",
    ])

    # ---- 8 ----
    doc.add_heading("8. 学习线（双轨，暂记）", level=1)
    add_bullets(doc, [
        "比赛线（Rule + Planner 保底）与 RL / DL 学习线并行，不冲突。",
        "RL 课程 = 实验矩阵（baseline → BC → 噪声 → 对手池 → POMDP → self-play），"
        "每行对应一个核心概念；配套读自己仓库的 train_ppo.py / diag_collapse.py / probe_policy.py。",
        "CV / DL 课程 = 视觉 90 天计划；年一唯一新增 DL 研究项 = 敌人行为模型"
        "（数据免费生成、可直接升级 Planner、标准监督学习）。",
        "资源：Sutton & Barto（选读）、Spinning Up、CleanRL。",
    ])

    # ---- 9 ----
    doc.add_heading("9. 待确认清单（等实车 / 官方）", level=1)
    add_table(doc, [
        ["事项", "找谁", "阻塞什么"],
        ["机械数单（见下）", "机械组", "环境运动学、CD 曲线实验、sim2real 参数"],
        ["隧道：净空高度 / 通行宽度 / 通行方式；压低后高度 / 转换时间 / 通过失败率"
         "（建议先做隧道断面样件试过）", "官方答疑 + 机械", "“压低过隧道”可行性、跨区机动"],
        ["交互卡检测范围（整区 vs 卡半径）", "官方答疑", "激活判定几何、点内机动自由度"],
        ["控制度 / 占领进度 / 伤害事件是否在哨兵串口流", "《2027 通信协议》/ 电控",
         "CD 直读 vs 自估"],
        ["高地分工（两个高地控制区靠队友？）", "队伍内部", "角色范围、CD 曲线实验口径"],
        ["赛前角色参数合法性", "裁判答疑", "赛前战术板机制"],
    ])
    doc.add_paragraph("机械数单（分时机）：")
    add_table(doc, [
        ["时机", "要什么", "拿来干什么"],
        ["现在", "轮径、轴距 / 轮距、外形包络（宽 × 高）", "里程计、全向解算、隧道通行第一判断"],
        ["设计冻结时", "v_max / a_max / 旋转角速度 / 质量 / 质心", "路径代价、CD 曲线实验、爬坡/侧翻裕量"],
        ["装车后实测", "功率曲线、缓冲行为、定位误差、路径跟踪误差", "sim2real 参数"],
    ])
    doc.add_paragraph("注：场地几何模型与机器人模型目前均不存在。数单到货前，环境内相关几何 / "
                      "运动学一律用占位值并标注 [占位]，相关结论不得作为最终依据。")

    # ---- 10 ----
    doc.add_heading("10. 更新记录", level=1)
    add_bullets(doc, [
        "2026-10-16 v1：基于多轮讨论收敛生成；等待实车参数与官方答复后更新。",
        "2026-10-16 v2：§7 改为三阶段推进顺序（A 抽象战术层 / B 保真升级 / C 部署）；"
        "§2 补“仿真环境（含场地几何模型）”归口；§9 补占位说明。",
        "后续更新点：轮径 / 轴距 / 运动性能 / 功率实测；隧道实测数；裁判系统数据字典；"
        "场地 / 机器人模型建成（B1/B2）；CD 曲线实验结果（人设判决）。",
    ])

    return doc


def main() -> None:
    doc = build()
    doc.save(OUT)
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    print("saved:", OUT)


if __name__ == "__main__":
    main()
